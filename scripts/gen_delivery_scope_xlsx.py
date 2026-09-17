#!/usr/bin/env python3
"""Export delivery scope checklist Excel for customer sign-off."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "木薯加工厂ERP-交付功能核对清单.xlsx"

# domain -> {module: (path, desc, group, channel)}
ADMIN_MODULES: dict[str, dict[str, tuple[str, str, str, str]]] = {
    "采购管理": {
        "农户档案": ("/purchase/hub/farmers", "农户主数据、默认单价、溯源前缀", "农户与入场", "管理端"),
        "过磅收货": ("/purchase/hub/weigh", "过磅单查询、流程状态（待入厂/待入库/待结算）", "农户与入场", "管理端+App"),
        "过磅流程编排": ("/purchase/hub/flow-graphs", "入厂流程节点配置", "农户与入场", "管理端"),
        "过磅品种": ("/purchase/hub/varieties", "品种/等级价维护", "农户与入场", "管理端"),
        "溯源批号": ("/purchase/hub/trace-batches", "溯源码批次管理", "农户与入场", "管理端"),
        "农户结算": ("/purchase/hub/settlements", "原料款结算、已付/待付", "结算", "管理端"),
        "原料溯源": ("/purchase/hub/trace", "溯源链查询", "农户与入场", "管理端"),
        "来料质检": ("/purchase/hub/qcs", "来料质检记录", "农户与入场", "管理端+App"),
    },
    "库存管理": {
        "库存查询": ("/inventory/hub/balances", "按仓/品/批号查库存", "库存与判断", "管理端+App"),
        "仓管待入库": ("/inventory/hub/inbound", "待入库过磅单、推仓状态", "库存与判断", "管理端+App"),
        "箱码管理": ("/inventory/hub/boxes", "板码/箱码台账", "库存与判断", "管理端+App"),
        "出入库记录汇总": ("/inventory/hub/stock-txns", "库存流水", "流水", "管理端"),
        "可用量分析": ("/inventory/hub/availability", "可发/可用量", "库存与判断", "管理端"),
        "亏料预警": ("/inventory/hub/shortage", "缺料预警", "库存与判断", "管理端"),
        "过量预警": ("/inventory/hub/excess", "超储预警", "库存与判断", "管理端"),
        "在途量统计": ("/inventory/hub/in-transits", "在途库存", "库存与判断", "管理端"),
        "待用量统计": ("/inventory/hub/reservations", "预留/待用量", "库存与判断", "管理端"),
    },
    "生产管理": {
        "工序定义": ("/production/hub/processes", "工序主数据", "工艺与规则", "管理端"),
        "工艺流程": ("/production/hub/routings", "工艺路线、工序顺序", "工艺与规则", "管理端"),
        "产线班次": ("/production/hub/shifts", "班次授权（控制 App 过站）", "工艺与规则", "管理端"),
        "例外派岗": ("/production/hub/dispatches", "灵活/例外派工", "工艺与规则", "管理端+App"),
        "工序流水": ("/production/hub/reports", "过站/报工记录", "现场台账", "管理端"),
        "计件工资": ("/production/hub/piecework", "计件汇总查询", "现场台账", "管理端"),
        "工序在制": ("/production/hub/process-wip", "在制品查询", "现场与溯源", "管理端"),
        "溯源生产": ("/production/hub/trace-production", "溯源批生产会话、启停结案", "现场与溯源", "管理端+App"),
        "工序扣损": ("/production/hub/process-yield", "工序收率/扣损", "现场与溯源", "管理端"),
        "退库未用完还仓": ("/production/hub/process-returns", "退库还仓记录", "现场台账", "管理端+App"),
    },
    "产品管理": {
        "产品档案": ("/product/hub/products", "原料/半成品/成品档案", "产品主数据", "管理端"),
        "产品单位管理": ("/product/hub/units", "计量单位", "产品主数据", "管理端"),
        "生产规格绑定": ("/product/hub/specs", "产品与工艺规格绑定", "产品主数据", "管理端"),
    },
    "工资管理": {
        "工人信息管理": ("/payroll/workers", "计件工人档案", "工价与档案", "管理端"),
        "工序工资": ("/payroll/wage-rates", "工序单价/工价", "工价与档案", "管理端"),
        "工资批量管理": ("/payroll/batch", "批量工资处理", "核算发放", "管理端"),
        "薪酬核算": ("/payroll/sheets", "月工资单", "核算发放", "管理端"),
        "员工工作台账": ("/payroll/work-records", "员工工作记录", "核算发放", "管理端"),
    },
    "人事管理": {
        "员工档案": ("/hr/employees", "员工主数据", "组织人事", "管理端"),
        "公司架构": ("/hr/departments", "部门/组织架构", "组织人事", "管理端"),
        "角色管理": ("/hr/roles", "角色与权限分配", "组织人事", "管理端"),
    },
    "财务管理": {
        "成本核算": ("/finance/hub/cost-accountings", "按期间/产品成本归集", "成本", "管理端"),
        "成本明细溯源表": ("/finance/hub/cost-traces", "成本来源追溯", "成本", "管理端"),
    },
    "统计报表": {
        "生产看板": ("/report/hub/production-board", "生产总览", "经营看板", "管理端"),
        "生产实况": ("/report/hub/live", "实时工序/溯源状态", "经营看板", "管理端"),
        "三仓库存概览": ("/report/hub/warehouse", "原料/半成品/成品仓", "经营看板", "管理端"),
        "日经营快照": ("/report/hub/daily", "日经营汇总", "日结对账", "管理端"),
        "原料入场日报": ("/report/hub/inbound-daily", "入场过磅日报", "日结对账", "管理端"),
        "计件日结汇总": ("/report/hub/piecework-daily", "计件日结", "日结对账", "管理端"),
        "工序扣损收率分析": ("/report/hub/yield-analysis", "收率分析", "分析查询", "管理端"),
        "收发存明细": ("/report/hub/stock-ledger", "库存收发存", "分析查询", "管理端"),
        "溯源批进度查询": ("/report/hub/trace-progress", "溯源批全流程进度", "分析查询", "管理端"),
        "农户结算对账汇总": ("/report/hub/farmer-settlement-summary", "农户款对账", "分析查询", "管理端"),
        "薪酬核算对账": ("/report/hub/payroll-reconcile", "工资单 vs 计件差异", "分析查询", "管理端"),
        "成本期间汇总": ("/report/hub/cost-period-summary", "期间成本汇总", "分析查询", "管理端"),
    },
    "系统管理": {
        "基础设置": ("/system/settings", "系统参数、载体码名称等", "基础与权限", "管理端"),
        "生产设置": ("/system/production-settings", "产线相关配置", "产线运维", "管理端"),
        "自定义权限": ("/iam/permissions", "权限码管理", "基础与权限", "管理端"),
        "登录控制": ("/iam/login-policy", "登录策略", "基础与权限", "管理端"),
        "批量核算工资": ("/system/batch-payroll-jobs", "批量工资任务", "产线运维", "管理端"),
        "操作日志": ("/automation/logs", "审计/操作日志", "产线运维", "管理端"),
    },
}

APP_MODULES = [
    ("员工端App", "生产（过站）", "/station", "选溯源批→扫工牌+板码→领料/退库/入库；溯源生产台", "计件工/固定工/班组长/管理员"),
    ("员工端App", "采购（收货）", "/receiving", "过磅入厂建单、过磅入库、出码推仓", "采购员/过磅员/管理员"),
    ("员工端App", "仓管作业", "/warehouse", "待入库、扫码核对入库、板码、库存、盘点", "仓管员/管理员"),
    ("员工端App", "班组管理", "/workshop", "任务派工、灵活派发、质检/返修/废料", "班组长/管理员"),
    ("员工端App", "质检工单", "QcShell", "待办工单、质检判定、履历", "质检员"),
    ("员工端App", "我的", "Tab", "今日计件/工钱、打卡、假勤、审批、工资、工牌", "全部角色"),
]

EXCLUDED = [
    ("销售域", "销售订单、客户 CRM、询价报价、销售外勤"),
    ("资产域", "固定资产全模块"),
    ("财务扩展", "完整总账（科目/凭证/发票/月结/三表等）；本期仅成本核算 2 项"),
    ("生产扩展", "MRP、自动 BOM、委外/受托、多单整合、进度跟踪"),
    ("客户侧", "客户自助 Web 门户"),
    ("平台级", "多租户 SaaS、复杂告警规则引擎"),
]

HEADER = ["序号", "终端", "业务域", "菜单分组", "功能模块", "访问路径", "功能说明", "适用角色/分工", "交付范围", "客户确认", "备注"]

THIN = Side(style="thin", color="C8D4D2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR_FILL = PatternFill("solid", fgColor="E8F4F2")
HDR_FONT = Font(name="微软雅黑", bold=True, size=10, color="0D7A6F")
BODY_FONT = Font(name="微软雅黑", size=10)
TITLE_FONT = Font(name="微软雅黑", bold=True, size=14, color="0D7A6F")
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_header(ws, row: int, ncol: int):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.border = BORDER
        cell.alignment = CENTER


def style_body(ws, r1: int, r2: int, ncol: int):
    for r in range(r1, r2 + 1):
        for c in range(1, ncol + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = WRAP if c not in (1, 9, 10) else CENTER


def set_widths(ws, widths: list[float]):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def sheet_overview(wb: Workbook):
    ws = wb.active
    ws.title = "交付总览"
    today = date.today().strftime("%Y年%m月%d日")
    ws.merge_cells("A1:F1")
    ws["A1"] = "木薯加工厂 ERP · 交付功能核对清单"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    info = [
        ["编制日期", today],
        ["版本", "木薯产线交付版（CASSAVA_PRODUCT_SCOPE）"],
        ["适用终端", "管理端 Web Admin + 员工端 Flutter App"],
        ["管理端功能模块", "58 项"],
        ["员工端现场模块", "6 项（含质检工单）"],
        ["统计报表", "12 项"],
        ["核心业务闭环", "入厂过磅 → 分板入库 → 溯源生产 → 计件核对 → 结算 → 报表"],
    ]
    r = 3
    for k, v in info:
        ws.cell(r, 1, k).font = Font(name="微软雅黑", bold=True, size=10)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        ws.cell(r, 2, v).font = BODY_FONT
        r += 1

    r += 1
    ws.cell(r, 1, "域统计").font = Font(name="微软雅黑", bold=True, size=11, color="0D7A6F")
    r += 1
    stats_hdr = ["业务域", "模块数量", "说明"]
    for c, h in enumerate(stats_hdr, 1):
        ws.cell(r, c, h)
    style_header(ws, r, 3)
    r += 1
    domain_notes = {
        "采购管理": "农户、过磅、溯源、结算",
        "库存管理": "库存、板码、预警、流水",
        "生产管理": "工艺、过站、溯源生产、计件",
        "产品管理": "产品主数据",
        "工资管理": "工价、薪酬核算",
        "人事管理": "员工、组织、角色",
        "财务管理": "成本核算（精简版）",
        "统计报表": "看板、日报、对账分析",
        "系统管理": "权限、设置、日志",
    }
    for domain, mods in ADMIN_MODULES.items():
        ws.append([domain, len(mods), domain_notes.get(domain, "")])
    ws.append(["员工端App", len(APP_MODULES), "现场过磅、仓管、过站、班组、质检、我的"])
    style_body(ws, r, r + len(ADMIN_MODULES), 3)
    set_widths(ws, [16, 12, 40])
    ws.freeze_panes = "A3"


def sheet_admin(wb: Workbook):
    ws = wb.create_sheet("管理端功能清单")
    ws.append(HEADER)
    style_header(ws, 1, len(HEADER))
    seq = 0
    row = 2
    for domain, mods in ADMIN_MODULES.items():
        for module, (path, desc, group, channel) in mods.items():
            seq += 1
            scope = "本期交付"
            confirm = ""
            ws.append([seq, "管理端", domain, group, module, path, desc, channel, scope, confirm, ""])
            row += 1
    style_body(ws, 2, row - 1, len(HEADER))
    set_widths(ws, [6, 10, 12, 12, 16, 28, 36, 14, 10, 10, 14])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}{row - 1}"


def sheet_app(wb: Workbook):
    ws = wb.create_sheet("员工端App功能")
    hdr = ["序号", "终端", "功能模块", "路由/入口", "功能说明", "适用角色", "交付范围", "客户确认", "备注"]
    ws.append(hdr)
    style_header(ws, 1, len(hdr))
    for i, (term, mod, route, desc, roles) in enumerate(APP_MODULES, 1):
        ws.append([i, term, mod, route, desc, roles, "本期交付", "", ""])
    style_body(ws, 2, 1 + len(APP_MODULES), len(hdr))
    set_widths(ws, [6, 12, 14, 14, 40, 22, 10, 10, 14])
    ws.freeze_panes = "A2"


def sheet_excluded(wb: Workbook):
    ws = wb.create_sheet("本期不包含")
    ws.append(["序号", "类别", "说明", "客户确认"])
    style_header(ws, 1, 4)
    for i, (cat, desc) in enumerate(EXCLUDED, 1):
        ws.append([i, cat, desc, ""])
    style_body(ws, 2, 1 + len(EXCLUDED), 4)
    set_widths(ws, [6, 14, 60, 12])
    ws.freeze_panes = "A2"


def sheet_signoff(wb: Workbook):
    ws = wb.create_sheet("签字确认")
    ws.merge_cells("A1:D1")
    ws["A1"] = "验收签字页"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = CENTER
    lines = [
        "",
        "经双方核对，确认《管理端功能清单》《员工端App功能》所列功能为木薯加工厂 ERP 本期交付范围。",
        "《本期不包含》所列功能不在本期合同交付范围内。",
        "",
        "甲方（客户）",
        "单位名称：",
        "授权代表：",
        "签字：                    日期：",
        "",
        "乙方（交付方）",
        "单位名称：",
        "授权代表：",
        "签字：                    日期：",
    ]
    for i, t in enumerate(lines, 3):
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=4)
        ws.cell(i, 1, t).font = BODY_FONT
        ws.cell(i, 1).alignment = Alignment(wrap_text=True, vertical="top")
    set_widths(ws, [20, 20, 20, 20])


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    sheet_overview(wb)
    sheet_admin(wb)
    sheet_app(wb)
    sheet_excluded(wb)
    sheet_signoff(wb)
    wb.save(OUT)
    print(f"OK: {OUT}")
    print(f"SIZE: {OUT.stat().st_size} bytes")


if __name__ == "__main__":
    main()
